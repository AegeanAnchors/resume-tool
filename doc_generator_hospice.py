"""
doc_generator_hospice.py

Builds the formatted Word document for the Hospice (RN & LPN) template.

Mirrors Home Health exactly — same base template, same layout rules:
  - "Organization" label instead of facility name header
  - Has Role field
  - Has Additional Info field (always shown for hospice jobs — yellow prompt if not found)
  - No Bed Count, Trauma Level, or Teaching Status fields
  - Patient Caseload is per day/week (not per shift)
  - Leadership/Charge appears ABOVE Additional Info
  - EMR/Charting and Patient Caseload only shown for is_hospice=True jobs
"""

import os
import shutil
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "Acute Formatted Example.docx")

YELLOW_PLACEHOLDER     = "[ Recruiter: Please Verify ]"
YELLOW_CREDENTIALS     = "[ Recruiter: Please Verify State Licenses & Certifications ]"
YELLOW_ADDITIONAL_INFO = "[ Recruiter: Please Verify Additional Hospice Info ]"
YELLOW_DATE            = "[ Date ]"


def generate_document_hospice(data: dict, output_path: str) -> str:
    """
    Builds the Hospice formatted Word document.
    Returns the output_path on success.
    """
    shutil.copy2(TEMPLATE_PATH, output_path)
    doc = Document(output_path)

    _clear_document_body(doc)

    candidate = data.get("candidate", {})
    education  = candidate.get("education") or []
    experience = data.get("experience") or []
    clinical   = data.get("clinical_experience") or []

    _add_candidate_name(doc, candidate)
    _add_blank_line(doc)
    _add_blank_line(doc)

    if education:
        _add_education_section(doc, education)
        _add_blank_line(doc)
        _add_blank_line(doc)

    _add_credentials_section(doc, candidate)
    _add_blank_line(doc)
    _add_blank_line(doc)

    _add_experience_section(doc, experience)

    if clinical:
        _add_blank_line(doc)
        _add_blank_line(doc)
        _add_clinical_experience_section(doc, clinical)

    phone = candidate.get("phone_number")
    _update_header_phone(doc, phone)
    _remove_trailing_blank_paragraphs(doc)

    doc.save(output_path)
    return output_path


# ─────────────────────────────────────────────
#  SECTION BUILDERS
# ─────────────────────────────────────────────

def _add_candidate_name(doc, candidate):
    name = _safe(candidate.get("full_name"), "UNKNOWN CANDIDATE").upper()
    credential = _safe(candidate.get("credential"))
    full_line = f"{name}, {credential}" if credential else name
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_space_after(para, 0)
    run = para.add_run(full_line)
    _set_run_font(run, size=14, bold=True, italic=True)


def _add_education_section(doc, education):
    _add_section_header(doc, "EDUCATION:")
    for edu in education:
        dates    = _safe(edu.get("dates"))
        school   = _safe(edu.get("school")).upper()
        location = _safe(edu.get("location"))
        school_line = f"{school}, {location}" if location else school

        # Education date + school — yellow date if missing, 2.5" column alignment
        _add_education_line(doc, dates, school_line)

        degree = _safe(edu.get("degree"))
        if degree:
            _add_indented_line(doc, degree, indent=3600)

        _add_blank_line(doc)


def _add_credentials_section(doc, candidate):
    _add_section_header(doc, "CREDENTIALS:")
    state_licenses = candidate.get("state_licenses") or []
    certifications = candidate.get("certifications") or []
    all_credentials = state_licenses + certifications

    if not all_credentials:
        para = doc.add_paragraph(style="EXP Bullet")
        _set_space_after(para, 0)
        _add_yellow_run(para, YELLOW_CREDENTIALS)
        return

    for item in all_credentials:
        para = doc.add_paragraph(style="EXP Bullet")
        _set_space_after(para, 0)
        run = para.add_run(_fix_credential_caps(item))
        _set_run_font(run, size=12, bold=False, italic=False)


def _add_experience_section(doc, experience):
    _add_section_header(doc, "EXPERIENCE:")
    for job in experience:
        _add_job_entry(doc, job)
        _add_blank_line(doc)


def _add_job_entry(doc, job):
    """
    Adds one complete Hospice job entry.

    Layout:
      MM/YYYY – MM/YYYY    ORGANIZATION NAME, City, ST
                           Employed Through: [value]
                           Discipline: RN (Staff)
                           Specialty: [value]
                           Role: [value or yellow]
                           EMR/Charting: [value or yellow]     ← hospice jobs only
                           Patient Caseload: [value or yellow] ← hospice jobs only
                           Leadership/Charge Experience: Yes   ← only if on resume
                           Additional Info: [value or yellow]
                           • Bullet points                     ← only if on resume
    """
    start_index = len(doc.paragraphs)

    dates   = _safe(job.get("dates"))
    org     = _safe(job.get("organization_name")).upper()
    city_st = _safe(job.get("organization_city_state"))
    org_line = f"{org}, {city_st}" if city_st else org

    # ── Date + Organization line ──────────────────────────────────────────
    _add_hanging_line(doc, f"{dates}\t{org_line}")

    # ── Employed Through ──────────────────────────────────────────────────
    _add_field_line(doc, "Employed Through", job.get("employed_through"))

    # ── Discipline: RN (Staff) ────────────────────────────────────────────
    discipline = _safe(job.get("discipline"))
    emp_type   = _safe(job.get("employment_type"))
    discipline_value = f"{discipline} ({emp_type})" if discipline and emp_type else discipline or None
    _add_field_line(doc, "Discipline", discipline_value)

    # ── Specialty ─────────────────────────────────────────────────────────
    _add_field_line(doc, "Specialty", job.get("specialty"))

    # ── Role ──────────────────────────────────────────────────────────────
    _add_field_line(doc, "Role", job.get("role"))

    is_hospice = job.get("is_hospice", True)

    # ── EMR/Charting — only for hospice jobs ──────────────────────────────
    if is_hospice:
        _add_field_line(doc, "EMR/Charting", job.get("emr"))

    # ── Patient Caseload — only for hospice jobs ──────────────────────────
    if is_hospice:
        _add_field_line(doc, "Patient Caseload", job.get("patient_caseload"))

    # ── Leadership/Charge Experience (only if on resume for this job) ─────
    if job.get("leadership_charge"):
        _add_field_line(doc, "Leadership/Charge Experience", "Yes")

    # ── Additional Info ───────────────────────────────────────────────────
    # Hospice jobs: bullet any hospice-specific items found; yellow prompt if none
    # Non-hospice jobs: bullet applicable items; skip section if none found
    additional_bullets = [b for b in (job.get("additional_info") or []) if b]

    if additional_bullets:
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        label_run = para.add_run("Additional Info:")
        _set_run_font(label_run, size=12, bold=False, italic=False)
        for item in additional_bullets:
            para = doc.add_paragraph(style="EXP Bullet")
            _set_space_after(para, 0)
            run = para.add_run(_safe(item))
            _set_run_font(run, size=12, bold=False, italic=False)
    elif is_hospice:
        # Hospice job with nothing found — always show yellow prompt
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        label_run = para.add_run("Additional Info: ")
        _set_run_font(label_run, size=12, bold=False, italic=False)
        _add_yellow_run(para, YELLOW_ADDITIONAL_INFO)
    # Non-hospice job with nothing found — skip Additional Info entirely

    # ── Bullet points (can break across pages) ───────────────────────────
    fields_end_index = len(doc.paragraphs)

    bullets = [b for b in (job.get("bullets") or []) if b]
    if bullets:
        for bullet_text in bullets:
            para = doc.add_paragraph(style="EXP Bullet")
            _set_space_after(para, 0)
            run = para.add_run(_safe(bullet_text))
            _set_run_font(run, size=12, bold=False, italic=False)

    # ── Lock structural fields together, let bullets flow freely ──────────
    for i in range(start_index, fields_end_index):
        doc.paragraphs[i].paragraph_format.keep_with_next = True


def _add_clinical_experience_section(doc, clinical):
    _add_section_header(doc, "CLINICAL EXPERIENCE:")
    for entry in clinical:
        facility = _safe(entry.get("facility_name")).upper()
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        run = para.add_run(facility)
        _set_run_font(run, size=12, bold=False, italic=False)

        rotation = _safe(entry.get("specialty_rotation")) or None
        if rotation:
            para2 = doc.add_paragraph(style="Position")
            _set_space_after(para2, 0)
            run2 = para2.add_run(rotation)
            _set_run_font(run2, size=12, bold=False, italic=False)

        _add_blank_line(doc)


# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────

def _safe(value, default=""):
    if value is None:
        return default
    return str(value).strip() or default


def _add_hanging_line(doc, text, tab_pos=2160):
    """
    Adds a line where the date sits at the left margin (0") and content
    after the tab starts at `tab_pos` twips. Wrapped text also stays at
    that same position.

    tab_pos defaults to 2160 (1.5") for experience dates (MM/YYYY format).
    Pass 3600 (2.5") for education dates written as full month names.
    """
    para = doc.add_paragraph()
    _set_space_after(para, 0)

    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(tab_pos))
    ind.set(qn("w:hanging"), str(tab_pos))
    pPr.append(ind)

    tabs_el = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "left")
    tab_el.set(qn("w:pos"), str(tab_pos))
    tabs_el.append(tab_el)
    pPr.append(tabs_el)

    run = para.add_run(text)
    _set_run_font(run, size=12, bold=False, italic=False)
    return para


def _add_indented_line(doc, text, indent=2160):
    """
    Adds a line that starts at exactly `indent` twips and wraps at the same point.
    Pass 3600 (2.5") for education degree lines.
    """
    para = doc.add_paragraph()
    _set_space_after(para, 0)

    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(indent))
    pPr.append(ind)

    run = para.add_run(text)
    _set_run_font(run, size=12, bold=False, italic=False)
    return para


def _add_education_line(doc, dates, school_line):
    """
    Education date + school name line with a 2.5" hanging indent.
    Two separate runs so the date can be highlighted yellow when missing.
    """
    para = doc.add_paragraph()
    _set_space_after(para, 0)

    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "3600")
    ind.set(qn("w:hanging"), "3600")
    pPr.append(ind)

    tabs_el = OxmlElement("w:tabs")
    tab_el = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "left")
    tab_el.set(qn("w:pos"), "3600")
    tabs_el.append(tab_el)
    pPr.append(tabs_el)

    if dates:
        date_run = para.add_run(dates)
        _set_run_font(date_run, size=12, bold=False, italic=False)
    else:
        date_run = para.add_run(YELLOW_DATE)
        _set_run_font(date_run, size=12, bold=False, italic=False)
        date_run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    school_run = para.add_run(f"\t{school_line}")
    _set_run_font(school_run, size=12, bold=False, italic=False)
    return para


def _add_section_header(doc, text):
    para = doc.add_paragraph()
    _set_space_after(para, 0)
    run = para.add_run(text)
    _set_run_font(run, size=12, bold=True, italic=False)


def _add_field_line(doc, label, value):
    para = doc.add_paragraph(style="Position")
    _set_space_after(para, 0)
    label_run = para.add_run(f"{label}: ")
    _set_run_font(label_run, size=12, bold=False, italic=False)
    if value:
        value_run = para.add_run(str(value))
        _set_run_font(value_run, size=12, bold=False, italic=False)
    else:
        _add_yellow_run(para, YELLOW_PLACEHOLDER)


def _add_yellow_run(para, text):
    run = para.add_run(text)
    _set_run_font(run, size=12, bold=False, italic=False)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _add_blank_line(doc):
    para = doc.add_paragraph()
    _set_space_after(para, 0)


def _set_space_after(para, points):
    para.paragraph_format.space_after = Pt(points)


def _set_run_font(run, size, bold, italic):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def _fix_credential_caps(text: str) -> str:
    """Ensures 'rn' and 'lpn' are always uppercase in credential bullets."""
    import re
    text = re.sub(r'\brn\b', 'RN', text, flags=re.IGNORECASE)
    text = re.sub(r'\blpn\b', 'LPN', text, flags=re.IGNORECASE)
    return text


def _update_header_phone(doc, phone_number):
    import re
    # Strip +1 country code prefix if present
    if phone_number:
        phone_number = re.sub(r'^\+1[\s\-\.]?', '', phone_number).strip()
    try:
        first_page_header = doc.sections[0].first_page_header
        for para in first_page_header.paragraphs:
            for run in para.runs:
                if re.search(r'\(?\d{3}\)?[\s\-\.]?\d{3}[\s\-\.]\d{4}', run.text):
                    run.text = phone_number if phone_number else "[ Recruiter: Please Verify Phone ]"
                    if not phone_number:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    return
    except Exception:
        pass


def _remove_trailing_blank_paragraphs(doc):
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    children = [c for c in list(body) if c != sect_pr]
    for child in reversed(children):
        if child.tag == qn("w:p"):
            text = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
            if text == "":
                body.remove(child)
            else:
                pPr = child.find(qn("w:pPr"))
                if pPr is not None:
                    keep_next = pPr.find(qn("w:keepNext"))
                    if keep_next is not None:
                        pPr.remove(keep_next)
                break
        else:
            break


def _clear_document_body(doc):
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)
