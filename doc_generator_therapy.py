"""
doc_generator_therapy.py

Builds the formatted Word document for the Therapy template
(PT, OT, SLP, PTA, COTA, RRT).
Uses Acute Formatted Example.docx as the base document.

Work experience entry fields:
  Dates + Facility (or Worksite for travel), Discipline (employment_type),
  Setting/Specialty, bullet points

Clinical Experience section:
  Dates + Facility, Role (Duration), Setting/Specialty, bullet points
"""

import os
import re
import shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "Acute Formatted Example.docx")

YELLOW_PLACEHOLDER = "[ Recruiter: Please Verify ]"
YELLOW_CREDENTIALS = "[ Recruiter: Please Verify State Licenses & Certifications ]"
YELLOW_DATE        = "[ Date ]"


# ─────────────────────────────────────────────
#  MAIN ENTRY POINT
# ─────────────────────────────────────────────

def generate_document_therapy(data: dict, output_path: str) -> str:
    shutil.copy2(TEMPLATE_PATH, output_path)
    doc = Document(output_path)

    _clear_document_body(doc)

    candidate  = data.get("candidate", {})
    education  = candidate.get("education") or []
    experience = data.get("experience") or []
    clinical   = data.get("clinical_experience") or []

    _add_candidate_name(doc, candidate)
    _add_blank_line(doc)
    _add_blank_line(doc)

    if education:
        _add_education_section(doc, education)
        _add_blank_line(doc)
        _add_blank_line(doc)

    _add_credentials_section(doc, candidate)
    _add_blank_line(doc)
    _add_blank_line(doc)

    if experience:
        _add_experience_section(doc, experience)

    if clinical:
        if experience:
            _add_blank_line(doc)
            _add_blank_line(doc)
        _add_clinical_experience_section(doc, clinical)

    _update_header_phone(doc, candidate.get("phone_number"))
    _remove_trailing_blank_paragraphs(doc)

    doc.save(output_path)
    return output_path


# ─────────────────────────────────────────────
#  SECTION BUILDERS
# ─────────────────────────────────────────────

def _add_candidate_name(doc, candidate):
    name       = _safe(candidate.get("full_name"), "UNKNOWN CANDIDATE").upper()
    credential = _safe(candidate.get("credential"))
    full_line  = f"{name}, {credential}" if credential else name
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_space_after(para, 0)
    run = para.add_run(full_line)
    _set_run_font(run, size=14, bold=True, italic=True)


def _add_education_section(doc, education):
    _add_section_header(doc, "EDUCATION:")
    for edu in education:
        dates    = _safe(edu.get("dates"))
        school   = _safe(edu.get("school")).upper()
        location = _safe(edu.get("location"))
        school_line = f"{school}, {location}" if location else school
        _add_education_line(doc, dates, school_line)
        degree = _safe(edu.get("degree"))
        if degree:
            _add_indented_line(doc, degree, indent=3600)
        _add_blank_line(doc)


def _add_credentials_section(doc, candidate):
    _add_section_header(doc, "CREDENTIALS:")
    profession     = _safe(candidate.get("profession"))
    state_licenses = candidate.get("state_licenses") or []
    certifications = candidate.get("certifications") or []

    # Therapy: profession first, then licenses, then certs
    all_credentials = []
    if profession:
        all_credentials.append(profession)
    all_credentials.extend(state_licenses)
    all_credentials.extend(certifications)

    if not all_credentials:
        para = doc.add_paragraph(style="EXP Bullet")
        _set_space_after(para, 0)
        _add_yellow_run(para, YELLOW_CREDENTIALS)
        return

    for item in all_credentials:
        para = doc.add_paragraph(style="EXP Bullet")
        _set_space_after(para, 0)
        run = para.add_run(_fix_therapy_caps(item))
        _set_run_font(run, size=12, bold=False, italic=False)


def _add_experience_section(doc, experience):
    _add_section_header(doc, "EXPERIENCE:")
    for job in experience:
        _add_job_entry(doc, job)
        _add_blank_line(doc)


def _add_job_entry(doc, job):
    """
    Therapy work experience entry:
      MM/YYYY – Present    [Worksite: ]FACILITY NAME, City, ST
                           Physical Therapist (Travel)
                           Outpatient Orthopedic
                           • Bullet points
    """
    start_index = len(doc.paragraphs)

    dates     = _safe(job.get("dates"))
    facility  = _safe(job.get("facility_name")).upper()
    city_st   = _safe(job.get("facility_city_state"))
    is_travel = job.get("is_travel", False)

    facility_display = f"Worksite: {facility}" if is_travel else facility
    fac_line = f"{facility_display}, {city_st}" if city_st else facility_display

    _add_hanging_line(doc, f"{dates}\t{fac_line}" if dates else f"\t{fac_line}")

    discipline = _safe(job.get("discipline"))
    emp_type   = _safe(job.get("employment_type"))
    if discipline and emp_type:
        disc_line = f"{discipline} ({emp_type})"
    else:
        disc_line = discipline or emp_type or None

    if disc_line:
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        run = para.add_run(disc_line)
        _set_run_font(run, size=12, bold=False, italic=False)

    setting = _safe(job.get("setting_specialty")) or None
    if setting:
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        run = para.add_run(setting)
        _set_run_font(run, size=12, bold=False, italic=False)
    else:
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        _add_yellow_run(para, YELLOW_PLACEHOLDER)

    fields_end_index = len(doc.paragraphs)

    for bullet_text in [b for b in (job.get("bullets") or []) if b]:
        para = doc.add_paragraph(style="EXP Bullet")
        _set_space_after(para, 0)
        run = para.add_run(_safe(bullet_text))
        _set_run_font(run, size=12, bold=False, italic=False)

    for i in range(start_index, fields_end_index):
        doc.paragraphs[i].paragraph_format.keep_with_next = True


def _add_clinical_experience_section(doc, clinical):
    _add_section_header(doc, "CLINICAL EXPERIENCE:")
    for entry in clinical:
        _add_clinical_entry(doc, entry)
        _add_blank_line(doc)


def _add_clinical_entry(doc, entry):
    """
    Clinical rotation entry:
      MM/YYYY – MM/YYYY    FACILITY NAME, City, ST
                           Physical Therapy Student (10 weeks)
                           Outpatient Sports Orthopedic
                           • Bullet points
    """
    start_index = len(doc.paragraphs)

    dates    = _safe(entry.get("dates"))
    facility = _safe(entry.get("facility_name")).upper()
    city_st  = _safe(entry.get("facility_city_state"))
    fac_line = f"{facility}, {city_st}" if city_st else facility

    _add_hanging_line(doc, f"{dates}\t{fac_line}" if dates else f"\t{fac_line}")

    role     = _safe(entry.get("role"))
    duration = _safe(entry.get("duration"))
    if role and duration:
        role_line = f"{role} ({duration})"
    elif role:
        role_line = role
    else:
        role_line = None

    if role_line:
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        run = para.add_run(role_line)
        _set_run_font(run, size=12, bold=False, italic=False)
    else:
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        _add_yellow_run(para, YELLOW_PLACEHOLDER)

    setting = _safe(entry.get("setting_specialty")) or None
    if setting:
        para = doc.add_paragraph(style="Position")
        _set_space_after(para, 0)
        run = para.add_run(setting)
        _set_run_font(run, size=12, bold=False, italic=False)

    fields_end_index = len(doc.paragraphs)

    for bullet_text in [b for b in (entry.get("bullets") or []) if b]:
        para = doc.add_paragraph(style="EXP Bullet")
        _set_space_after(para, 0)
        run = para.add_run(_safe(bullet_text))
        _set_run_font(run, size=12, bold=False, italic=False)

    for i in range(start_index, fields_end_index):
        doc.paragraphs[i].paragraph_format.keep_with_next = True


# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────

def _safe(value, default=""):
    if value is None:
        return default
    return str(value).strip() or default


def _fix_therapy_caps(text):
    for abbr in ["PT", "OT", "SLP", "PTA", "COTA", "RRT", "CRT",
                 "DPT", "MPT", "OTD", "MOT", "OTR", "BLS", "ACLS",
                 "RN", "LPN"]:
        text = re.sub(rf'\b{abbr}\b', abbr, text, flags=re.IGNORECASE)
    return text


def _add_hanging_line(doc, text, tab_pos=2160):
    para = doc.add_paragraph()
    _set_space_after(para, 0)
    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"),    str(tab_pos))
    ind.set(qn("w:hanging"), str(tab_pos))
    pPr.append(ind)
    tabs_el = OxmlElement("w:tabs")
    tab_el  = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "left")
    tab_el.set(qn("w:pos"), str(tab_pos))
    tabs_el.append(tab_el)
    pPr.append(tabs_el)
    run = para.add_run(text)
    _set_run_font(run, size=12, bold=False, italic=False)
    return para


def _add_indented_line(doc, text, indent=2160):
    para = doc.add_paragraph()
    _set_space_after(para, 0)
    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(indent))
    pPr.append(ind)
    run = para.add_run(text)
    _set_run_font(run, size=12, bold=False, italic=False)
    return para


def _add_education_line(doc, dates, school_line):
    para = doc.add_paragraph()
    _set_space_after(para, 0)
    pPr = para._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"),    "3600")
    ind.set(qn("w:hanging"), "3600")
    pPr.append(ind)
    tabs_el = OxmlElement("w:tabs")
    tab_el  = OxmlElement("w:tab")
    tab_el.set(qn("w:val"), "left")
    tab_el.set(qn("w:pos"), "3600")
    tabs_el.append(tab_el)
    pPr.append(tabs_el)
    if dates:
        date_run = para.add_run(dates)
        _set_run_font(date_run, size=12, bold=False, italic=False)
    else:
        date_run = para.add_run(YELLOW_DATE)
        _set_run_font(date_run, size=12, bold=False, italic=False)
        date_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    school_run = para.add_run(f"\t{school_line}")
    _set_run_font(school_run, size=12, bold=False, italic=False)
    return para


def _add_section_header(doc, text):
    para = doc.add_paragraph()
    _set_space_after(para, 0)
    run = para.add_run(text)
    _set_run_font(run, size=12, bold=True, italic=False)


def _add_yellow_run(para, text):
    run = para.add_run(text)
    _set_run_font(run, size=12, bold=False, italic=False)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _add_blank_line(doc):
    para = doc.add_paragraph()
    _set_space_after(para, 0)


def _set_space_after(para, points):
    para.paragraph_format.space_after = Pt(points)


def _set_run_font(run, size, bold, italic):
    run.font.name   = "Times New Roman"
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic


def _update_header_phone(doc, phone_number):
    if phone_number:
        phone_number = re.sub(r'^\+1[\s\-\.]?', '', phone_number).strip()
    try:
        first_page_header = doc.sections[0].first_page_header
        for para in first_page_header.paragraphs:
            for run in para.runs:
                if re.search(r'\(?\d{3}\)?[\s\-\.]?\d{3}[\s\-\.]\d{4}', run.text):
                    run.text = phone_number if phone_number else "[ Recruiter: Please Verify Phone ]"
                    if not phone_number:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    return
    except Exception:
        pass


def _remove_trailing_blank_paragraphs(doc):
    body    = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    children = [c for c in list(body) if c != sect_pr]
    for child in reversed(children):
        if child.tag == qn("w:p"):
            text = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
            if text == "":
                body.remove(child)
            else:
                pPr = child.find(qn("w:pPr"))
                if pPr is not None:
                    keep_next = pPr.find(qn("w:keepNext"))
                    if keep_next is not None:
                        pPr.remove(keep_next)
                break
        else:
            break


def _clear_document_body(doc):
    body    = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)
